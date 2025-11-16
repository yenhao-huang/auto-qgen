#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel Parser
將 Excel 檔案解析為 JSON 格式，包含每個儲存格的詳細資訊
"""

import argparse
import json
import os
import sys
from pathlib import Path
from typing import List, Dict, Any, Union
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet


def get_column_name(col_index: int) -> str:
    """
    將欄位索引轉換為欄位名稱 (A, B, C, ..., AA, AB, ...)

    Args:
        col_index: 欄位索引 (1-based)

    Returns:
        欄位名稱字串
    """
    column_name = ""
    while col_index > 0:
        col_index -= 1
        column_name = chr(col_index % 26 + ord('A')) + column_name
        col_index //= 26
    return column_name


def parse_excel_sheet(sheet: Worksheet, sheet_name: str) -> List[Dict[str, Any]]:
    """
    解析單一工作表

    Args:
        sheet: openpyxl 工作表物件
        sheet_name: 工作表名稱

    Returns:
        包含所有儲存格資訊的列表
    """
    cells_data = []

    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:  # 只記錄有內容的儲存格
                cell_info = {
                    "sheet_name": sheet_name,
                    "row_name": str(cell.row),
                    "col_name": get_column_name(cell.column),
                    "cell_address": f"{get_column_name(cell.column)}{cell.row}",
                    "content": cell.value
                }
                cells_data.append(cell_info)

    return cells_data


def parse_excel_file(file_path: Union[str, Path], include_empty_cells: bool = False) -> Dict[str, Any]:
    """
    解析單一 Excel 檔案

    Args:
        file_path: Excel 檔案路徑
        include_empty_cells: 是否包含空白儲存格 (預設為 False)

    Returns:
        包含檔案資訊和所有儲存格資料的字典
    """
    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(f"檔案不存在: {file_path}")

    if file_path.suffix not in ['.xlsx', '.xlsm', '.xltx', '.xltm']:
        raise ValueError(f"不支援的檔案格式: {file_path.suffix}")

    workbook = openpyxl.load_workbook(file_path, data_only=True)

    result = {
        "file_name": file_path.name,
        "file_path": str(file_path.absolute()),
        "sheets": [],
        "total_cells": 0
    }

    all_cells = []

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]

        if include_empty_cells:
            # 包含所有儲存格
            sheet_cells = []
            for row in sheet.iter_rows():
                for cell in row:
                    cell_info = {
                        "sheet_name": sheet_name,
                        "row_name": str(cell.row),
                        "col_name": get_column_name(cell.column),
                        "cell_address": f"{get_column_name(cell.column)}{cell.row}",
                        "content": cell.value
                    }
                    sheet_cells.append(cell_info)
        else:
            # 只包含有內容的儲存格
            sheet_cells = parse_excel_sheet(sheet, sheet_name)

        all_cells.extend(sheet_cells)

        result["sheets"].append({
            "sheet_name": sheet_name,
            "cell_count": len(sheet_cells)
        })

    result["cells"] = all_cells
    result["total_cells"] = len(all_cells)

    return result


def parse_excel_directory(dir_path: Union[str, Path],
                         output_dir: Union[str, Path] = None,
                         include_empty_cells: bool = False) -> List[Dict[str, Any]]:
    """
    解析目錄中的所有 Excel 檔案

    Args:
        dir_path: 目錄路徑
        output_dir: 輸出 JSON 檔案的目錄 (可選)
        include_empty_cells: 是否包含空白儲存格

    Returns:
        所有檔案的解析結果列表
    """
    dir_path = Path(dir_path)

    if not dir_path.exists():
        raise FileNotFoundError(f"目錄不存在: {dir_path}")

    if not dir_path.is_dir():
        raise NotADirectoryError(f"不是目錄: {dir_path}")

    # 尋找所有 Excel 檔案
    excel_files = list(dir_path.glob("*.xlsx")) + \
                  list(dir_path.glob("*.xlsm")) + \
                  list(dir_path.glob("*.xltx")) + \
                  list(dir_path.glob("*.xltm"))

    results = []

    for excel_file in excel_files:
        try:
            print(f"解析檔案: {excel_file.name}")
            result = parse_excel_file(excel_file, include_empty_cells)
            results.append(result)

            # 如果指定輸出目錄，則儲存個別 JSON 檔案
            if output_dir:
                output_path = Path(output_dir)
                output_path.mkdir(parents=True, exist_ok=True)

                json_filename = excel_file.stem + ".json"
                json_path = output_path / json_filename

                with open(json_path, 'w', encoding='utf-8') as f:
                    json.dump(result, f, ensure_ascii=False, indent=2)
                print(f"  已儲存: {json_path}")

        except Exception as e:
            print(f"  錯誤: {e}")
            continue

    return results


def save_to_json(data: Union[Dict, List], output_path: Union[str, Path]) -> None:
    """
    將資料儲存為 JSON 檔案

    Args:
        data: 要儲存的資料
        output_path: 輸出檔案路徑
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"已儲存 JSON 檔案: {output_path}")


def main():
    """
    主程式 - CLI 介面
    """
    parser = argparse.ArgumentParser(
        description="Excel Parser - 將 Excel 檔案解析為 JSON 格式",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
使用範例:
  # 解析單一檔案（使用預設檔案名稱）
  python excel_parser.py

  # 解析指定的單一檔案
  python excel_parser.py --file 表12-1.xlsx -o output.json

  # 解析整個目錄（使用預設目錄）
  python excel_parser.py --mode dir

  # 解析指定目錄中的所有 Excel 檔案
  python excel_parser.py --mode dir --dir raw_data/一般服務 -o parsed_data

  # 包含空白儲存格
  python excel_parser.py --file 表12-1.xlsx --include-empty

  # 顯示摘要資訊
  python excel_parser.py --file 表12-1.xlsx --summary-only
        """
    )

    parser.add_argument(
        "--mode",
        choices=["file", "dir"],
        default="file",
        help="處理模式：file (單一檔案) 或 dir (整個目錄)，預設為 file"
    )

    parser.add_argument(
        "--file",
        "-f",
        type=str,
        default="raw_data/一般服務/表12-1.xlsx",
        help="要解析的 Excel 檔案路徑（預設：表12-1.xlsx）"
    )

    parser.add_argument(
        "--dir",
        "-d",
        type=str,
        default="raw_data/一般服務",
        help="要解析的目錄路徑（預設：raw_data/一般服務）"
    )

    parser.add_argument(
        "--output",
        "-o",
        type=str,
        default=None,
        help="輸出檔案或目錄路徑（預設：file 模式輸出到 raw_data/parsed_data/{檔名}.json，dir 模式輸出到 raw_data/parsed_data/）"
    )

    parser.add_argument(
        "--include-empty",
        action="store_true",
        help="包含空白儲存格（預設：否）"
    )

    parser.add_argument(
        "--summary-only",
        action="store_true",
        help="僅顯示摘要資訊，不儲存完整 JSON（預設：否）"
    )

    parser.add_argument(
        "--quiet",
        "-q",
        action="store_true",
        help="安靜模式，減少輸出訊息（預設：否）"
    )

    args = parser.parse_args()

    if not args.quiet:
        print("=" * 70)
        print("Excel Parser")
        print("=" * 70)

    try:
        if args.mode == "file":
            # 單一檔案模式
            file_path = Path(args.file)

            if not file_path.exists():
                print(f"錯誤：檔案不存在 - {file_path}", file=sys.stderr)
                sys.exit(1)

            if not args.quiet:
                print(f"\n解析檔案: {file_path}")

            result = parse_excel_file(file_path, args.include_empty)

            # 顯示摘要
            if not args.quiet:
                print(f"  檔案名稱: {result['file_name']}")
                print(f"  工作表數: {len(result['sheets'])}")
                for sheet in result['sheets']:
                    print(f"    - {sheet['sheet_name']}: {sheet['cell_count']} 個儲存格")
                print(f"  總儲存格數: {result['total_cells']}")

            # 儲存結果
            if not args.summary_only:
                if args.output:
                    output_path = Path(args.output)
                else:
                    output_path = Path("raw_data/parsed_data") / f"{file_path.stem}.json"

                save_to_json(result, output_path)

                if not args.quiet:
                    print(f"\n✓ 已儲存至: {output_path}")

        elif args.mode == "dir":
            # 目錄模式
            dir_path = Path(args.dir)

            if not dir_path.exists():
                print(f"錯誤：目錄不存在 - {dir_path}", file=sys.stderr)
                sys.exit(1)

            if not dir_path.is_dir():
                print(f"錯誤：路徑不是目錄 - {dir_path}", file=sys.stderr)
                sys.exit(1)

            if not args.quiet:
                print(f"\n解析目錄: {dir_path}")

            # 設定輸出目錄
            if args.summary_only:
                output_dir = None
            else:
                output_dir = args.output if args.output else "raw_data/parsed_data"

            results = parse_excel_directory(dir_path, output_dir, args.include_empty)

            # 顯示摘要
            if not args.quiet or args.summary_only:
                print(f"\n總共解析 {len(results)} 個檔案:")
                for r in results:
                    print(f"  - {r['file_name']}: {r['total_cells']} 個儲存格")

            # 儲存總覽
            if not args.summary_only and output_dir:
                summary = {
                    "total_files": len(results),
                    "files": [
                        {
                            "file_name": r["file_name"],
                            "sheets": len(r["sheets"]),
                            "total_cells": r["total_cells"]
                        }
                        for r in results
                    ]
                }

                summary_path = Path(output_dir) / "summary.json"
                save_to_json(summary, summary_path)

                if not args.quiet:
                    print(f"\n✓ 總覽已儲存至: {summary_path}")

        if not args.quiet:
            print("\n" + "=" * 70)
            print("✓ 完成！")
            print("=" * 70)

    except FileNotFoundError as e:
        print(f"錯誤：{e}", file=sys.stderr)
        sys.exit(1)
    except ValueError as e:
        print(f"錯誤：{e}", file=sys.stderr)
        sys.exit(1)
    except Exception as e:
        print(f"未預期的錯誤：{e}", file=sys.stderr)
        if not args.quiet:
            import traceback
            traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    main()
